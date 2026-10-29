# -*- coding: utf-8 -*-
"""
Hisob raqami ko'chirmasi tahlili: hisobga tushib, qaytib chiqmagan tushumlarni aniqlash.

Kirish : bank tizimidan olingan "Справка о работе счета" hisoboti (.xlsb / .xlsx / .xls)
Chiqish: Excel fayl (Xulosa, Qaytmagan tushumlar, To'lovchilar kesimida,
         Bog'lanmagan chiqimlar, Moslangan juftliklar)

Ishlatish:
    python hisob_tahlil.py                      -> oyna (GUI) ochiladi
    python hisob_tahlil.py fayl.xlsb [natija.xlsx]  -> buyruq satri rejimi
"""
import os
import re
import sys
import itertools
import traceback
from collections import defaultdict
from datetime import datetime

import pandas as pd

VERSION = "1.0"

# --------------------------------------------------------------------------- #
#  Matn yordamchi funksiyalari
# --------------------------------------------------------------------------- #
_TR = dict(zip(
    'АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯЎҚҒҲ',
    ['A', 'B', 'V', 'G', 'D', 'E', 'YO', 'J', 'Z', 'I', 'Y', 'K', 'L', 'M', 'N', 'O', 'P', 'R', 'S', 'T',
     'U', 'F', 'X', 'TS', 'CH', 'SH', 'SH', '', 'I', '', 'E', 'YU', 'YA', 'O', 'Q', 'G', 'H']))


def translit(s):
    """Kirill -> lotin (taqqoslash uchun soddalashtirilgan)."""
    return ''.join(_TR.get(c, c) for c in str(s).upper())


def norm(s):
    """Taqqoslash uchun normallashtirish: faqat A-Z0-9, Q->K, H->X."""
    s = translit(s).replace('Q', 'K').replace('H', 'X')
    return re.sub(r'[^A-Z0-9]', '', s)


_ID_RE = re.compile(r'(?<!\d)(\d{12,15})(?!\d)')          # ijro hujjati / PINFL / ish raqami
_CASE_RE = re.compile(r'\d-\d{4}-\d{4}/\d+(?:-\d+)?')       # sud ishi raqami 2-1001-2405/4316-2
_PINFL_RE = re.compile(r'(?<!\d)([3-6]\d{13})(?!\d)')
_STOP = {'YATT', 'YTT', 'CHP', 'OOO', 'MCHJ', 'IP'}


def extract_ids(s):
    out = set(_ID_RE.findall(s)) | set(_CASE_RE.findall(s))
    # hisob raqamlari bo'laklarini chiqarib tashlaymiz
    return frozenset(x for x in out if not x.startswith('2980') and not x.startswith('2020'))


def surname_from_receiver(name):
    words = [w for w in re.findall(r'[A-Z]+', translit(name)) if len(w) > 3 and w not in _STOP]
    return norm(words[0]) if words else ''


def parse_balance(a, p):
    """'0.00 (А)', '66 979 089.16 (П)' -> passiv qoldiq (П - А)."""
    def num(x):
        x = re.sub(r'\(.*?\)', '', str(x)).replace(' ', '').replace('\xa0', '').replace(',', '.')
        try:
            return float(x)
        except ValueError:
            return 0.0
    return num(p) - num(a)


DATE_RE = re.compile(r'^\d{2}\.\d{2}\.\d{4}$')


def to_date(v):
    if isinstance(v, (int, float)) and not pd.isna(v):        # Excel serial sana
        return pd.Timestamp('1899-12-30') + pd.Timedelta(days=float(v))
    s = str(v).strip()
    if DATE_RE.match(s):
        return pd.to_datetime(s, format='%d.%m.%Y')
    return None


# --------------------------------------------------------------------------- #
#  1. Faylni o'qish
# --------------------------------------------------------------------------- #
COLS = ['date', 'doc', 'vo', 'acc_payer', 'name_payer', 'mfo_payer',
        'acc_rcv', 'name_rcv', 'mfo_rcv', 'dt', 'kt', 'sum', 'purpose']


def read_statement(path, log=print):
    ext = os.path.splitext(path)[1].lower()
    engine = {'.xlsb': 'pyxlsb', '.xls': 'xlrd'}.get(ext, 'openpyxl')
    log(f"Fayl o'qilmoqda: {os.path.basename(path)}")
    raw = pd.read_excel(path, engine=engine, header=None, dtype=object)
    if raw.shape[1] < 13:
        raise ValueError("Fayl tuzilishi kutilganidek emas (kamida 13 ustun bo'lishi kerak).")
    raw = raw.iloc[:, :13]
    log(f"  {raw.shape[0]} qator topildi")

    info = {'account': '', 'bank': '', 'period': '', 'opening': None, 'closing': None}
    first = str(raw.iat[0, 0])
    m = re.search(r'№\s*(\d{20})', first)
    if m:
        info['account'] = m.group(1)
    m = re.search(r'с\s+(\d{2}\.\d{2}\.\d{4})\s+по\s+(\d{2}\.\d{2}\.\d{4})', first)
    if m:
        info['period'] = f"{m.group(1)} – {m.group(2)}"
    info['bank'] = str(raw.iat[1, 0]).strip() if raw.shape[0] > 1 else ''

    # Sarlavha qismidagi kirish / chiqish qoldig'i (birinchi uchragani)
    for i in range(min(40, len(raw))):
        c0 = str(raw.iat[i, 0])
        if info['opening'] is None and c0.startswith('Входящий остаток') and not pd.isna(raw.iat[i, 5]):
            info['opening'] = parse_balance(raw.iat[i, 4], raw.iat[i, 5])
        if info['closing'] is None and c0.startswith('Исходящий остаток') and not pd.isna(raw.iat[i, 5]):
            info['closing'] = parse_balance(raw.iat[i, 4], raw.iat[i, 5])

    # Operatsiyalarni ajratish ("Дата текущей операции" bloklari bo'yicha)
    rows, op = [], None
    for i in range(len(raw)):
        c0 = raw.iat[i, 0]
        s0 = str(c0)
        m = re.search(r'Дата текущей операции\s+(\d{2}\.\d{2}\.\d{4})', s0)
        if m:
            op = pd.to_datetime(m.group(1), format='%d.%m.%Y')
            continue
        d = to_date(c0)
        if d is None:
            continue
        r = dict(zip(COLS, raw.iloc[i].tolist()))
        r['date'] = d
        r['op_date'] = op if op is not None else d
        r['row'] = i + 1                                    # Excel qatori raqami
        rows.append(r)
    tx = pd.DataFrame(rows)
    if tx.empty:
        raise ValueError("Faylda operatsiyalar topilmadi.")
    for c in ['dt', 'kt', 'sum']:
        tx[c] = pd.to_numeric(tx[c], errors='coerce').fillna(0.0)
    for c in ['purpose', 'name_rcv', 'name_payer', 'acc_payer', 'acc_rcv', 'doc', 'vo', 'mfo_payer', 'mfo_rcv']:
        tx[c] = tx[c].map(lambda v: '' if pd.isna(v) else str(v).strip())
    log(f"  Operatsiyalar: {len(tx)} (Kt: {(tx.kt > 0).sum()}, Dt: {(tx.dt > 0).sum()})")
    if not info['account']:
        # hisob raqami sarlavhada bo'lmasa — eng ko'p uchragan hisob
        info['account'] = pd.concat([tx.acc_payer, tx.acc_rcv]).mode().iat[0]
    return tx, info


# --------------------------------------------------------------------------- #
#  2. Moslashtirish (tushum <-> chiqim)
# --------------------------------------------------------------------------- #
STAGE_NAMES = {
    1: "Summa + to'lov matni bir xil",
    2: "Summa + ijro hujjati / PINFL / ish raqami",
    3: "Summa + familiya",
    4: "Summa teng, 31 kun ichida",
    5: "Bir nechta tushum = bitta chiqim",
    6: "Summa ~ teng (≤0,5%) + shaxs",
    7: "Storno (chiqim qaytib tushgan)",
    8: "Bitta tushum → ko'p chiqim",
    9: "FIFO yig'indi (familiya bo'yicha)",
}


def match(tx, log=print):
    tx = tx.copy()
    tx['ids'] = tx.purpose.map(extract_ids)
    tx['npurp'] = tx.purpose.map(norm)
    tx['sur'] = tx.name_rcv.map(surname_from_receiver)
    tx['d'] = tx.op_date.values.astype('datetime64[D]').astype('int64')

    C = tx[tx.kt > 0].to_dict('records')
    B = tx[tx.dt > 0].sort_values(['op_date', 'row'], kind='stable').to_dict('records')
    for r in C:
        r['amt'] = round(r['kt'], 2)
    for r in B:
        r['amt'] = round(r['dt'], 2)

    by_amt = defaultdict(list)
    for i, c in enumerate(C):
        by_amt[c['amt']].append(i)
    mc, mb = {}, {}   # tushum i -> (chiqim k, bosqich);  chiqim k -> [tushumlar]

    def score(c, d):
        s, p = 0, c['npurp']
        if len(p) > 15 and p[:60] in d['npurp']:
            s += 100
        if c['ids'] & d['ids']:
            s += 50
        if len(d['sur']) > 3 and d['sur'] in p:
            s += 30
        return s

    # 1-4 bosqich: summa bo'yicha birga-bir
    for stage, minscore, maxdays in [(1, 100, 99999), (2, 50, 99999), (3, 30, 99999), (4, 0, 31)]:
        for k, d in enumerate(B):
            if k in mb:
                continue
            cands = [i for i in by_amt.get(d['amt'], [])
                     if i not in mc and 0 <= d['d'] - C[i]['d'] <= maxdays]
            if not cands:
                continue
            best = max(cands, key=lambda i: (score(C[i], d), C[i]['d']))
            if score(C[best], d) >= minscore:
                mc[best] = (k, stage)
                mb[k] = [best]
        log(f"  {stage}-bosqich tugadi: moslangan tushumlar {len(mc)}")

    # 5-bosqich: bir nechta tushum yig'indisi = bitta chiqim
    id_idx = defaultdict(set)
    for i, c in enumerate(C):
        for x in c['ids']:
            id_idx[x].add(i)
    for k, d in enumerate(B):
        if k in mb:
            continue
        pool = set()
        for x in d['ids']:
            pool |= id_idx.get(x, set())
        pool = [i for i in pool if i not in mc and C[i]['d'] <= d['d']]
        if sum(C[i]['amt'] for i in pool) + 0.01 < d['amt'] and len(d['sur']) > 3:
            pool = [i for i, c in enumerate(C) if i not in mc and c['d'] <= d['d'] and d['sur'] in c['npurp']]
        if not pool:
            continue
        pool = sorted(pool, key=lambda i: -C[i]['d'])[:18]
        found = None
        if abs(sum(C[i]['amt'] for i in pool) - d['amt']) < 0.02:
            found = pool
        else:
            for n in range(2, min(5, len(pool)) + 1):
                for comb in itertools.combinations(pool, n):
                    if abs(sum(C[i]['amt'] for i in comb) - d['amt']) < 0.02:
                        found = list(comb)
                        break
                if found:
                    break
        if found:
            for i in found:
                mc[i] = (k, 5)
            mb[k] = found
    log(f"  5-bosqich tugadi: {len(mc)}")

    # 6-bosqich: summa deyarli teng + shaxs bir xil
    for k, d in enumerate(B):
        if k in mb:
            continue
        tol = max(1000.0, d['amt'] * 0.005)
        cands = [i for i, c in enumerate(C) if i not in mc and 0 <= d['d'] - c['d'] <= 400
                 and abs(c['amt'] - d['amt']) <= tol and score(c, d) >= 30]
        if cands:
            best = max(cands, key=lambda i: (score(C[i], d), -abs(C[i]['amt'] - d['amt'])))
            mc[best] = (k, 6)
            mb[k] = [best]
    log(f"  6-bosqich tugadi: {len(mc)}")

    # 7-bosqich: storno — chiqim keyinroq hisobga qaytib tushgan
    for k, d in enumerate(B):
        if k in mb:
            continue
        cands = [i for i in by_amt.get(d['amt'], []) if i not in mc and 0 <= C[i]['d'] - d['d'] <= 120
                 and (C[i]['acc_payer'] == d['acc_rcv']
                      or (len(d['sur']) > 3 and d['sur'] in norm(C[i]['name_payer'] + C[i]['purpose'])))]
        if cands:
            best = min(cands, key=lambda i: C[i]['d'])
            mc[best] = (k, 7)
            mb[k] = [best]
    log(f"  7-bosqich tugadi: {len(mc)}")

    # 8-bosqich: bitta tushum bir kunda bir xil matn bilan ko'p hisobga taqsimlangan
    grp = defaultdict(list)
    for k, d in enumerate(B):
        if k not in mb:
            grp[(d['d'], d['purpose'].strip().upper())].append(k)
    for (day, _), ks in grp.items():
        if len(ks) < 2:
            continue
        tot = sum(B[k]['amt'] for k in ks)
        cands = [i for i, c in enumerate(C) if i not in mc and 0 <= day - c['d'] <= 31
                 and abs(c['amt'] - tot) <= max(1.0, tot * 0.01)]
        if cands:
            best = min(cands, key=lambda i: abs(C[i]['amt'] - tot))
            mc[best] = (ks[0], 8)
            for k in ks:
                mb[k] = [best]
    log(f"  8-bosqich tugadi: {len(mc)}")

    # 9-bosqich: chiqim = shu shaxsning avvalgi tushumlari FIFO yig'indisi
    for k, d in enumerate(B):
        if k in mb or len(d['sur']) <= 4:
            continue
        pool = sorted([i for i, c in enumerate(C) if i not in mc and c['d'] <= d['d'] and d['sur'] in c['npurp']],
                      key=lambda i: C[i]['d'])
        acc, take = 0.0, []
        for i in pool:
            acc += C[i]['amt']
            take.append(i)
            if abs(acc - d['amt']) <= max(1.0, d['amt'] * 0.005):
                for t in take:
                    mc[t] = (k, 9)
                mb[k] = take
                break
            if acc > d['amt']:
                break
    log(f"  9-bosqich tugadi: {len(mc)}")

    cr, db = pd.DataFrame(C), pd.DataFrame(B)
    cr['matched'] = [i in mc for i in range(len(C))]
    cr['stage'] = [mc[i][1] if i in mc else None for i in range(len(C))]
    cr['debit_row'] = [B[mc[i][0]]['row'] if i in mc else None for i in range(len(C))]
    db['matched'] = [k in mb for k in range(len(B))]
    db['credit_rows'] = [','.join(str(C[i]['row']) for i in mb[k]) if k in mb else '' for k in range(len(B))]
    return cr, db


# --------------------------------------------------------------------------- #
#  3. Natijani Excelga yozish
# --------------------------------------------------------------------------- #
def export(cr, db, info, out_path, log=print):
    uc = cr[~cr.matched].copy()
    ud = db[~db.matched].copy()

    def related(c_npurp, c_ids, d_sur, d_ids):
        return (len(d_sur) > 4 and d_sur in c_npurp) or bool(c_ids & d_ids)

    ud_list = ud[['row', 'd', 'sur', 'ids']].to_dict('records')
    uc_list = uc[['row', 'd', 'npurp', 'ids']].to_dict('records')
    uc['rel'] = [','.join(str(d['row']) for d in ud_list
                          if d['d'] >= c['d'] and related(c['npurp'], c['ids'], d['sur'], d['ids'])) for c in uc_list]
    ud['rel'] = [','.join(str(c['row']) for c in uc_list
                          if c['d'] <= d['d'] and related(c['npurp'], c['ids'], d['sur'], d['ids'])) for d in ud_list]

    uc['PINFL'] = uc.purpose.map(lambda s: ', '.join(sorted(set(_PINFL_RE.findall(s)))))
    ref_day = cr.op_date.max()
    uc['days'] = (ref_day - uc.op_date).dt.days
    uc['cat'] = pd.cut(uc.kt, [0, 1e5, 1e6, 1e7, 1e8, 1e15],
                       labels=['≤100 ming', '100 ming – 1 mln', '1 – 10 mln', '10 – 100 mln', '>100 mln'])

    opening = info.get('opening')
    closing = info.get('closing')
    has_rel = uc.rel != ''
    summary = pd.DataFrame([
        ['Hisob raqami', info.get('account', ''), None],
        ['Bank / filial', info.get('bank', ''), None],
        ['Davr', info.get('period', ''), None],
        ['Kirish qoldig‘i', opening, None],
        ['Jami tushumlar (Kt)', cr.kt.sum(), len(cr)],
        ['Jami chiqimlar (Dt)', db.dt.sum(), len(db)],
        ['Chiqish qoldig‘i', closing, None],
        ['Chiqim bilan moslangan tushumlar', cr[cr.matched].kt.sum(), int(cr.matched.sum())],
        ['QAYTIB CHIQMAGAN TUSHUMLAR', uc.kt.sum(), len(uc)],
        ['   shundan: shu shaxsga qisman chiqim bo‘lgan (tekshirish kerak)', uc[has_rel].kt.sum(), int(has_rel.sum())],
        ['   shundan: hech qanday chiqim izi yo‘q', uc[~has_rel].kt.sum(), int((~has_rel).sum())],
        ['Hech bir tushumga bog‘lanmagan chiqimlar', ud.dt.sum(), len(ud)],
        ['Tekshiruv: qaytmagan tushum − bog‘lanmagan chiqim', uc.kt.sum() - ud.dt.sum(), None],
        ['Tekshiruv: chiqish qoldig‘i − kirish qoldig‘i',
         (closing - opening) if (opening is not None and closing is not None) else None, None],
        ['Tahlil sanasi', datetime.now().strftime('%d.%m.%Y %H:%M'), None],
    ], columns=["Ko'rsatkich", 'Qiymat', 'Soni'])

    by_cat = uc.groupby('cat', observed=False).agg(Soni=('kt', 'size'), Summa=('kt', 'sum')).reset_index() \
        .rename(columns={'cat': 'Summa toifasi'})
    by_month = uc.groupby(uc.op_date.dt.strftime('%Y-%m')).agg(Soni=('kt', 'size'), Summa=('kt', 'sum')).reset_index() \
        .rename(columns={'op_date': 'Oy'})
    by_stage = cr[cr.matched].groupby('stage').agg(Soni=('kt', 'size'), Summa=('kt', 'sum')).reset_index()
    by_stage['stage'] = by_stage.stage.map(STAGE_NAMES)
    by_stage = by_stage.rename(columns={'stage': 'Moslash usuli'})
    by_payer = uc.groupby('name_payer').agg(Soni=('kt', 'size'), Summa=('kt', 'sum')) \
        .sort_values('Summa', ascending=False).reset_index().rename(columns={'name_payer': "To'lovchi"})

    cols_c = {'row': 'Excel qatori', 'op_date': 'Operatsiya sanasi', 'date': 'Hujjat sanasi', 'doc': 'Hujjat №',
              'vo': 'ВО', 'acc_payer': "To'lovchi hisobi", 'name_payer': "To'lovchi", 'mfo_payer': 'MFO',
              'kt': 'Summa (Kt)', 'PINFL': 'PINFL', 'purpose': "To'lov maqsadi", 'days': 'Necha kun turibdi',
              'cat': 'Summa toifasi', 'rel': "Bog'liq bo'lishi mumkin bo'lgan chiqim qatori"}
    cols_d = {'row': 'Excel qatori', 'op_date': 'Operatsiya sanasi', 'doc': 'Hujjat №', 'acc_rcv': 'Oluvchi hisobi',
              'name_rcv': 'Oluvchi', 'dt': 'Summa (Dt)', 'purpose': "To'lov maqsadi",
              'rel': "Bog'liq bo'lishi mumkin bo'lgan tushum qatori"}

    mp = cr[cr.matched].copy()
    mp['method'] = mp.stage.map(STAGE_NAMES)
    mp = mp.merge(db[['row', 'op_date', 'name_rcv', 'acc_rcv', 'dt']].rename(
        columns={'row': 'debit_row', 'op_date': 'd_date', 'name_rcv': 'd_name', 'acc_rcv': 'd_acc', 'dt': 'd_amt'}),
        on='debit_row', how='left')
    mp = mp[['row', 'op_date', 'name_payer', 'kt', 'purpose', 'method', 'debit_row', 'd_date', 'd_name', 'd_acc', 'd_amt']]
    mp.columns = ['Tushum qatori', 'Tushum sanasi', "To'lovchi", 'Tushum summasi', "To'lov maqsadi", 'Moslash usuli',
                  'Chiqim qatori', 'Chiqim sanasi', 'Oluvchi', 'Oluvchi hisobi', 'Chiqim summasi']

    log("Excel fayl yozilmoqda...")
    from openpyxl.styles import Font, PatternFill
    with pd.ExcelWriter(out_path, engine='openpyxl') as w:
        r = 0
        summary.to_excel(w, sheet_name='Xulosa', index=False, startrow=r); r += len(summary) + 3
        by_cat.to_excel(w, sheet_name='Xulosa', index=False, startrow=r); r += len(by_cat) + 3
        by_stage.to_excel(w, sheet_name='Xulosa', index=False, startrow=r); r += len(by_stage) + 3
        by_month.to_excel(w, sheet_name='Xulosa', index=False, startrow=r)
        uc.sort_values('kt', ascending=False)[list(cols_c)].rename(columns=cols_c) \
            .to_excel(w, sheet_name='Qaytmagan tushumlar', index=False)
        by_payer.to_excel(w, sheet_name="To'lovchilar kesimida", index=False)
        ud.sort_values('dt', ascending=False)[list(cols_d)].rename(columns=cols_d) \
            .to_excel(w, sheet_name="Bog'lanmagan chiqimlar", index=False)
        mp.to_excel(w, sheet_name='Moslangan juftliklar', index=False)

        hdr_fill = PatternFill('solid', fgColor='DDEBF7')
        for ws in w.book.worksheets:
            ws.freeze_panes = 'A2' if ws.title != 'Xulosa' else None
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
                for c in col:
                    if isinstance(c.value, float):
                        c.number_format = '#,##0.00'
                    elif isinstance(c.value, datetime):
                        c.number_format = 'DD.MM.YYYY'
            if ws.title != 'Xulosa':
                for c in ws[1]:
                    c.font = Font(bold=True)
                    c.fill = hdr_fill
                ws.auto_filter.ref = ws.dimensions
        xs = w.book['Xulosa']
        for row in xs.iter_rows():
            if row[0].value in ("Ko'rsatkich", 'Summa toifasi', 'Moslash usuli', 'Oy'):
                for c in row:
                    c.font = Font(bold=True)
                    c.fill = hdr_fill
            if row[0].value == 'QAYTIB CHIQMAGAN TUSHUMLAR':
                for c in row:
                    c.font = Font(bold=True, color='C00000')
    return summary, uc, ud


def run(path, out_path=None, log=print):
    tx, info = read_statement(path, log)
    log("Tushum va chiqimlar moslashtirilmoqda...")
    cr, db = match(tx, log)
    if not out_path:
        base = os.path.splitext(path)[0]
        out_path = f"{base}_{info.get('account', '')}_qaytmagan_tushumlar.xlsx"
    summary, uc, ud = export(cr, db, info, out_path, log)
    log("")
    log(f"Hisob: {info.get('account')}  |  Davr: {info.get('period')}")
    log(f"Qaytib chiqmagan tushumlar: {len(uc)} ta, {uc.kt.sum():,.2f} so'm".replace(',', ' '))
    log(f"Bog'lanmagan chiqimlar   : {len(ud)} ta, {ud.dt.sum():,.2f} so'm".replace(',', ' '))
    log(f"Natija: {out_path}")
    return out_path


# --------------------------------------------------------------------------- #
#  4. Oyna (GUI)
# --------------------------------------------------------------------------- #
def gui():
    import threading
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title(f"Hisob ko'chirmasi tahlili v{VERSION} — qaytmagan tushumlar")
    root.geometry('780x520')

    src = tk.StringVar()
    dst = tk.StringVar()

    def pick_src():
        p = filedialog.askopenfilename(title="Ko'chirma faylini tanlang",
                                       filetypes=[('Excel', '*.xlsb *.xlsx *.xls'), ('Barcha fayllar', '*.*')])
        if p:
            src.set(p)
            dst.set(os.path.splitext(p)[0] + '_qaytmagan_tushumlar.xlsx')

    def pick_dst():
        p = filedialog.asksaveasfilename(title='Natijani saqlash', defaultextension='.xlsx',
                                         filetypes=[('Excel', '*.xlsx')])
        if p:
            dst.set(p)

    frm = ttk.Frame(root, padding=10)
    frm.pack(fill='both', expand=True)
    ttk.Label(frm, text="Ko'chirma fayli (.xlsb/.xlsx):").grid(row=0, column=0, sticky='w')
    ttk.Entry(frm, textvariable=src, width=80).grid(row=1, column=0, sticky='we')
    ttk.Button(frm, text='Tanlash...', command=pick_src).grid(row=1, column=1, padx=5)
    ttk.Label(frm, text='Natija fayli (.xlsx):').grid(row=2, column=0, sticky='w', pady=(8, 0))
    ttk.Entry(frm, textvariable=dst, width=80).grid(row=3, column=0, sticky='we')
    ttk.Button(frm, text='Tanlash...', command=pick_dst).grid(row=3, column=1, padx=5)

    btn = ttk.Button(frm, text='▶  Tahlil qilish')
    btn.grid(row=4, column=0, sticky='w', pady=10)
    bar = ttk.Progressbar(frm, mode='indeterminate')
    bar.grid(row=4, column=0, sticky='e', pady=10, ipadx=100)

    txt = scrolledtext.ScrolledText(frm, height=20, font=('Consolas', 9))
    txt.grid(row=5, column=0, columnspan=2, sticky='nsew')
    frm.columnconfigure(0, weight=1)
    frm.rowconfigure(5, weight=1)

    def log(msg):
        root.after(0, lambda: (txt.insert('end', str(msg) + '\n'), txt.see('end')))

    def work():
        try:
            out = run(src.get(), dst.get() or None, log)
            root.after(0, lambda: done(out, None))
        except Exception as e:
            log(traceback.format_exc())
            root.after(0, lambda: done(None, e))

    def done(out, err):
        bar.stop()
        btn.config(state='normal')
        if err:
            messagebox.showerror('Xatolik', str(err))
        elif messagebox.askyesno('Tayyor', f"Natija saqlandi:\n{out}\n\nFaylni ochaymi?"):
            try:
                os.startfile(out)
            except Exception:
                pass

    def start():
        if not os.path.isfile(src.get()):
            messagebox.showwarning('Diqqat', "Avval ko'chirma faylini tanlang.")
            return
        txt.delete('1.0', 'end')
        btn.config(state='disabled')
        bar.start(10)
        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    root.mainloop()


if __name__ == '__main__':
    if len(sys.argv) > 1:
        run(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    else:
        gui()
